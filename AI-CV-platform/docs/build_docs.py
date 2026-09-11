#!/usr/bin/env python3
"""Generate the user manual (.docx) and 3-slide intro deck (.pptx)."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn as ppt_qn
from pptx.util import Emu, Inches, Pt as PptPt
from lxml import etree

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "戴纳AI引导式视觉平台-使用说明书.md"
DOCX_PATH = ROOT / "戴纳AI引导式视觉平台-使用说明书.docx"
PPTX_PATH = ROOT / "戴纳AI引导式视觉平台-产品介绍.pptx"

CN_FONT = "微软雅黑"
CN_FONT_FALLBACK = "WenQuanYi Micro Hei"

# Platform palette
NAVY = (10, 14, 21)
PANEL = (17, 24, 35)
PANEL2 = (22, 32, 46)
PANEL3 = (29, 40, 55)
LINE = (36, 49, 70)
INK = (232, 239, 249)
MUTE = (139, 160, 187)
BRAND = (56, 189, 248)
BRAND_DIM = (14, 165, 233)
OK = (52, 211, 153)
WARN = (251, 191, 36)
WHITE = (255, 255, 255)
DARK_INK = (29, 40, 55)
BODY = (50, 62, 78)


def rgb(doc_color: tuple[int, int, int]) -> RGBColor:
    return RGBColor(*doc_color)


def set_east_asia_font(run, name: str = CN_FONT) -> None:
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:cs"), name)


def shade_cell(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:shd"):
            tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, color="C5D0DE") -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_runs(paragraph, text: str, size=11, color=BODY, bold=False) -> None:
    """Render a line that may contain **bold** and `code` spans."""
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        if part.startswith("**") and part.endswith("**"):
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run.text = part[1:-1]
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
            run.font.size = Pt(size - 0.5)
            run.font.color.rgb = rgb(BRAND_DIM)
            continue
        else:
            run.text = part
            run.bold = bold
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)
        set_east_asia_font(run)


def style_paragraph(p, space_before=0, space_after=8, line=1.25, align=None) -> None:
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align


def add_heading_styled(doc, text: str, level: int) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_east_asia_font(run)
    if level == 1:
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = rgb(BRAND_DIM)
        style_paragraph(p, space_before=18, space_after=8, line=1.2)
        # bottom border
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "0EA5E9")
        pBdr.append(bottom)
        pPr.append(pBdr)
    else:
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = rgb(DARK_INK)
        style_paragraph(p, space_before=12, space_after=6, line=1.2)


def add_code_block(doc, code: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, "F4F7FB")
    set_cell_border(cell, "D5DEEA")
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(code.rstrip() + "\n")
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = rgb(DARK_INK)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    style_paragraph(p, space_before=4, space_after=4, line=1.15)
    doc.add_paragraph()


def add_table(doc, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, "0EA5E9")
        set_cell_border(cell, "0EA5E9")
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = rgb(WHITE)
        set_east_asia_font(run)
        style_paragraph(p, space_before=3, space_after=3, line=1.15)
    for r, row in enumerate(rows):
        fill = "F0F7FC" if r % 2 == 0 else "FFFFFF"
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            shade_cell(cell, fill)
            set_cell_border(cell, "D5DEEA")
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, val, size=10, color=DARK_INK)
            style_paragraph(p, space_before=2, space_after=2, line=1.15)
    spacer = doc.add_paragraph()
    style_paragraph(spacer, space_before=2, space_after=8)


def set_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("戴纳 AI 引导式视觉平台  ·  使用说明书")
    run.font.size = Pt(9)
    run.font.color.rgb = rgb(MUTE)
    set_east_asia_font(run)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("v0.3.0-mvp    ·    ")
    run.font.size = Pt(9)
    run.font.color.rgb = rgb(MUTE)
    set_east_asia_font(run)

    # PAGE field
    run2 = fp.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run2._r.append(fld_begin)
    run2._r.append(instr)
    run2._r.append(fld_sep)
    t = OxmlElement("w:t")
    t.text = "1"
    run2._r.append(t)
    run2._r.append(fld_end)
    run2.font.size = Pt(9)
    run2.font.color.rgb = rgb(MUTE)


def add_cover(doc: Document) -> None:
    for _ in range(3):
        p = doc.add_paragraph()
        style_paragraph(p, space_after=0)

    bar = doc.add_table(rows=1, cols=1)
    cell = bar.cell(0, 0)
    shade_cell(cell, "0EA5E9")
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run("DINAR  ·  INDUSTRIAL VISION")
    run.font.size = Pt(11)
    run.font.color.rgb = rgb(WHITE)
    run.bold = True
    set_east_asia_font(run)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_before=10, space_after=10)

    p = doc.add_paragraph()
    style_paragraph(p, space_before=36, space_after=6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("戴纳 AI 引导式视觉平台")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = rgb(DARK_INK)
    set_east_asia_font(run)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("使  用  说  明  书")
    run.font.size = Pt(20)
    run.font.color.rgb = rgb(BRAND_DIM)
    set_east_asia_font(run)
    style_paragraph(p, space_before=8, space_after=18)

    meta = [
        "软件版本    v0.3.0-mvp",
        "文档日期    2026 年 8 月 26 日",
        "适用对象    视觉工程师  /  现场操作员  /  项目负责人",
    ]
    for line in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.font.size = Pt(12)
        run.font.color.rgb = rgb(BODY)
        set_east_asia_font(run)
        style_paragraph(p, space_after=4)

    p = doc.add_paragraph()
    style_paragraph(p, space_before=28, space_after=8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("用一句话描述检测需求，平台自动生成可运行、可发布的视觉方案。")
    run.font.size = Pt(12)
    run.italic = True
    run.font.color.rgb = rgb(MUTE)
    set_east_asia_font(run)

    doc.add_page_break()


def parse_table(lines: list[str], i: int) -> tuple[list[str], list[list[str]], int]:
    header = [c.strip() for c in lines[i].strip().strip("|").split("|")]
    i += 1
    if i < len(lines) and re.match(r"^\s*\|?\s*-+", lines[i]):
        i += 1
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
        i += 1
    return header, rows, i


def build_docx() -> None:
    md = MD_PATH.read_text(encoding="utf-8")
    lines = md.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(0.8)

    styles = doc.styles["Normal"]
    styles.font.name = CN_FONT
    styles.font.size = Pt(11)
    styles._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    add_cover(doc)
    set_header_footer(doc)

    toc = doc.add_paragraph()
    run = toc.add_run("目录")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = rgb(BRAND_DIM)
    set_east_asia_font(run)
    style_paragraph(toc, space_after=10)

    toc_items = [
        "1. 产品概述",
        "2. 安装与启动",
        "3. 界面与角色",
        "4. 项目管理",
        "5. 数据与标注",
        "6. AI 助手",
        "7. 视觉流程",
        "8. 批量测试",
        "9. 模型训练",
        "10. 发布与运行",
        "11. 监控",
        "12. 设备",
        "13. 推荐作业流程",
        "14. 算子一览",
        "15. 常见问题",
        "16. 部署边界与后续",
        "附录 A  快捷对照",
        "附录 B  术语",
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        run = p.add_run(item)
        run.font.size = Pt(12)
        run.font.color.rgb = rgb(DARK_INK)
        set_east_asia_font(run)
        style_paragraph(p, space_after=4, line=1.4)
    doc.add_page_break()

    i = 0
    in_code = False
    code_buf: list[str] = []
    skipped_title = False

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if line.startswith("```"):
            if in_code:
                add_code_block(doc, "\n".join(code_buf))
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(raw)
            i += 1
            continue

        if not line:
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            continue
        if line.startswith("# ") and not skipped_title:
            skipped_title = True
            i += 1
            # skip subtitle / meta block under H1 until first ##
            while i < len(lines) and not lines[i].startswith("## "):
                i += 1
            continue
        if line.startswith("## "):
            add_heading_styled(doc, line[3:].strip(), 1)
            i += 1
            continue
        if line.startswith("### "):
            add_heading_styled(doc, line[4:].strip(), 2)
            i += 1
            continue
        if line.strip().startswith("|"):
            headers, rows, i = parse_table(lines, i)
            add_table(doc, headers, rows)
            continue

        # lists
        m_ul = re.match(r"^(\s*)[-*] (.+)$", line)
        m_ol = re.match(r"^(\s*)(\d+)\. (.+)$", line)
        if m_ul or m_ol:
            text = m_ul.group(2) if m_ul else m_ol.group(3)
            indent = len((m_ul.group(1) if m_ul else m_ol.group(1)) or "")
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6 + indent * 0.25)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            bullet = "• " if m_ul else f"{m_ol.group(2)}. "
            run = p.add_run(bullet)
            run.font.size = Pt(11)
            run.font.color.rgb = rgb(BRAND_DIM)
            set_east_asia_font(run)
            add_runs(p, text, size=11, color=BODY)
            style_paragraph(p, space_after=3, line=1.3)
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, line, size=11, color=BODY)
        style_paragraph(p, space_after=8, line=1.35)
        i += 1

    doc.save(DOCX_PATH)
    print(f"wrote {DOCX_PATH}")


# ---------- PPT ----------

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def ppt_rgb(c: tuple[int, int, int]) -> PptRGB:
    return PptRGB(*c)


def set_run_font(run, name=CN_FONT, size=14, bold=False, color=INK) -> None:
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = ppt_rgb(color)
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(ppt_qn(tag))
        if el is None:
            el = etree.SubElement(rPr, ppt_qn(tag))
        el.set("typeface", name)


def add_rect(slide, l, t, w, h, fill, line=None, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, l, t, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = ppt_rgb(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = ppt_rgb(line)
        shp.line.width = Emu(6350)
    if radius is not None:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    shp.shadow.inherit = False
    return shp


def tb(slide, l, t, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor="t"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=bold, color=color)
    # anchor
    body = box._element.find(ppt_qn("p:txBody"))
    bodyPr = body.find(ppt_qn("a:bodyPr"))
    if bodyPr is not None:
        bodyPr.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
        bodyPr.set("lIns", "0")
        bodyPr.set("rIns", "0")
        bodyPr.set("tIns", "0")
        bodyPr.set("bIns", "0")
    return box


def tb_multi(slide, l, t, w, h, lines, anchor="t"):
    """lines: list of dicts with text/size/bold/color/align/space_after."""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    body = box._element.find(ppt_qn("p:txBody"))
    bodyPr = body.find(ppt_qn("a:bodyPr"))
    if bodyPr is not None:
        bodyPr.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
        bodyPr.set("lIns", "0")
        bodyPr.set("rIns", "0")
        bodyPr.set("tIns", "0")
        bodyPr.set("bIns", "0")
    for i, spec in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", PP_ALIGN.LEFT)
        p.space_after = PptPt(spec.get("space", 0))
        p.space_before = PptPt(spec.get("space_before", 0))
        run = p.add_run()
        run.text = spec["text"]
        set_run_font(
            run,
            size=spec.get("size", 14),
            bold=spec.get("bold", False),
            color=spec.get("color", INK),
        )
    return box


def paint_bg(slide, color=NAVY) -> None:
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = ppt_rgb(color)
    bg.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def chrome(slide, page: int, total: int = 3) -> None:
    add_rect(slide, 0, 0, Inches(0.12), SLIDE_H, BRAND)
    tb(slide, Inches(0.45), Inches(0.18), Inches(8), Inches(0.28),
       "戴纳 AI 引导式视觉平台", size=11, bold=True, color=BRAND)
    tb(slide, Inches(10.4), Inches(0.18), Inches(2.5), Inches(0.28),
       f"v0.3.0-mvp    {page} / {total}", size=11, color=MUTE, align=PP_ALIGN.RIGHT)
    add_rect(slide, Inches(0.45), Inches(0.50), Inches(12.4), Emu(12700), LINE)


def build_pptx() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ===== SLIDE 1 =====
    s1 = prs.slides.add_slide(blank)
    paint_bg(s1)
    chrome(s1, 1)

    tb(s1, Inches(0.55), Inches(0.85), Inches(12), Inches(0.32),
       "工业视觉  ·  引导式开发平台", size=14, color=BRAND)

    tb(s1, Inches(0.55), Inches(1.22), Inches(12.2), Inches(0.7),
       "用一句话描述检测需求", size=36, bold=True, color=WHITE)
    tb(s1, Inches(0.55), Inches(1.90), Inches(12.2), Inches(0.55),
       "平台自动生成可运行、可发布的视觉方案", size=28, bold=True, color=BRAND)

    tb(s1, Inches(0.55), Inches(2.55), Inches(12.2), Inches(0.45),
       "传统图像算法与深度学习分类在同一条流程里完成：诊断成像、推荐路线、画布调参、批量验证漏检误报，发布后操作员只看到业务阈值和 OK / NG。",
       size=15, color=MUTE)

    pillars = [
        ("01", "说人话出方案", "意图识别 + 成像诊断\n生成白名单内可运行流程"),
        ("02", "画布上调得动", "66 个算子 DAG 执行\n缓存重算、逐节点预览"),
        ("03", "发布改不坏", "参数分工程师 / 业务 / 锁定\n操作员只能改范围内阈值"),
    ]
    for i, (num, title, desc) in enumerate(pillars):
        x = Inches(0.55 + i * 4.2)
        add_rect(s1, x, Inches(3.22), Inches(3.95), Inches(2.05), PANEL2, radius=0.08)
        add_rect(s1, x, Inches(3.22), Inches(0.10), Inches(2.05), BRAND)
        tb(s1, x + Inches(0.28), Inches(3.35), Inches(3.4), Inches(0.32),
           num, size=12, bold=True, color=BRAND)
        tb(s1, x + Inches(0.28), Inches(3.68), Inches(3.4), Inches(0.4),
           title, size=18, bold=True, color=WHITE)
        tb(s1, x + Inches(0.28), Inches(4.18), Inches(3.5), Inches(0.85),
           desc, size=13, color=MUTE)

    stats = [
        ("66", "视觉算子"),
        ("3", "分类训练后端"),
        ("2", "工作角色"),
        ("1", "条上线闭环"),
    ]
    add_rect(s1, Inches(0.0), Inches(5.55), SLIDE_W, Inches(1.95), PANEL)
    for i, (n, label) in enumerate(stats):
        x = Inches(0.7 + i * 3.2)
        tb(s1, x, Inches(5.85), Inches(2.6), Inches(0.55), n, size=28, bold=True, color=BRAND, align=PP_ALIGN.CENTER)
        tb(s1, x, Inches(6.42), Inches(2.6), Inches(0.35), label, size=13, color=MUTE, align=PP_ALIGN.CENTER)
        if i < 3:
            add_rect(s1, Inches(3.6 + i * 3.2), Inches(6.05), Emu(12700), Inches(0.7), LINE)

    # ===== SLIDE 2 =====
    s2 = prs.slides.add_slide(blank)
    paint_bg(s2)
    chrome(s2, 2)
    tb(s2, Inches(0.55), Inches(0.68), Inches(12), Inches(0.42),
       "从需求到量产检测，一条闭环", size=24, bold=True, color=WHITE)
    tb(s2, Inches(0.55), Inches(1.12), Inches(12), Inches(0.32),
       "工程师在画布上完成方案，操作员在运行页只面对结果。", size=13, color=MUTE)

    steps = ["需求", "数据", "AI 方案", "流程", "批量", "发布", "监控"]
    n = len(steps)
    left0 = Inches(0.45)
    gap = Inches(0.16)
    usable = Inches(12.4)
    step_w = int((usable - gap * (n - 1)) / n)
    y = Inches(1.58)
    h = Inches(0.62)
    for i, name in enumerate(steps):
        x = left0 + i * (step_w + gap)
        fill = BRAND_DIM if i in (0, n - 1) else PANEL3
        add_rect(s2, x, y, step_w, h, fill, radius=0.2)
        tb(s2, x, y, step_w, h, name, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor="m")
        if i < n - 1:
            tb(s2, x + step_w - Inches(0.04), y, gap + Inches(0.08), h, "›",
               size=18, bold=True, color=BRAND, align=PP_ALIGN.CENTER, anchor="m")

    cards = [
        ("AI 助手", "用业务语言描述目标。平台做意图识别与成像诊断，生成通过校验的节点流程；无大模型时走规则引擎。"),
        ("节点画布", "66 个算子按端口类型连线。改参自动重跑下游，缓存命中前段，ROI 可在图上框选回填。"),
        ("数据标注", "上传或目录导入。整图分类 / 框 / 多边形 / 点。导出 COCO、YOLO、分类清单。"),
        ("模型训练", "LBP+SVM 适合小样本纹理；EfficientNet 与 YOLO-cls 适合外观分类，训完导出 ONNX。"),
        ("批量测试", "准确率、精确率、召回率、F1、漏检、误报。误判可回流训练集，结果导出 CSV。"),
        ("运行监控", "发布后操作员只改业务参数。监控看设备在线、置信度、耗时、良率与逐次明细。"),
    ]
    for i, (title, desc) in enumerate(cards):
        col, row = i % 3, i // 3
        x = Inches(0.45 + col * 4.25)
        y = Inches(2.48 + row * 2.22)
        add_rect(s2, x, y, Inches(4.05), Inches(2.05), PANEL2, radius=0.07)
        add_rect(s2, x, y, Inches(4.05), Inches(0.08), BRAND)
        tb(s2, x + Inches(0.22), y + Inches(0.22), Inches(3.6), Inches(0.4),
           title, size=16, bold=True, color=WHITE)
        tb(s2, x + Inches(0.22), y + Inches(0.68), Inches(3.62), Inches(1.2),
           desc, size=13, color=MUTE)

    # ===== SLIDE 3 =====
    s3 = prs.slides.add_slide(blank)
    paint_bg(s3)
    chrome(s3, 3)
    tb(s3, Inches(0.55), Inches(0.68), Inches(12), Inches(0.42),
       "典型场景，两种角色，立刻可跑", size=24, bold=True, color=WHITE)

    scenes = [
        ("表面缺陷", "污渍 · 瑕疵 · 异物",
         "CLAHE、自适应阈值、形态学、Blob。面积或数量超限即 NG。适合布面、零件外观。"),
        ("物料分类", "袋装谷物 · 外观种类",
         "按类别导入样本，训 EfficientNet 或 YOLO 分类，接到 AI 分类节点。业务参数是最低置信度。"),
        ("生物视觉", "iPSC 克隆定位",
         "局部纹理分割克隆区域，输出面积占比、等效直径。也可分块分类做形态分型。"),
    ]
    for i, (title, tag, desc) in enumerate(scenes):
        x = Inches(0.45 + i * 4.25)
        add_rect(s3, x, Inches(1.25), Inches(4.05), Inches(2.55), PANEL2, radius=0.07)
        tb(s3, x + Inches(0.22), Inches(1.40), Inches(3.6), Inches(0.35),
           title, size=18, bold=True, color=WHITE)
        tb(s3, x + Inches(0.22), Inches(1.78), Inches(3.6), Inches(0.28),
           tag, size=12, color=BRAND)
        tb(s3, x + Inches(0.22), Inches(2.15), Inches(3.62), Inches(1.4),
           desc, size=13, color=MUTE)

    # dual role
    add_rect(s3, Inches(0.45), Inches(4.05), Inches(6.25), Inches(2.55), PANEL2, radius=0.07)
    add_rect(s3, Inches(0.45), Inches(4.05), Inches(0.10), Inches(2.55), BRAND)
    tb(s3, Inches(0.80), Inches(4.20), Inches(5.6), Inches(0.35),
       "工程师模式", size=16, bold=True, color=BRAND)
    for j, line in enumerate([
        "搭流程、看中间图、训模型、发版本",
        "全部节点参数可见，含工程师级阈值",
        "批量测试量化漏检 / 误报后再发布",
    ]):
        tb(s3, Inches(0.80), Inches(4.65 + j * 0.42), Inches(5.6), Inches(0.38),
           "▸  " + line, size=13, color=INK)

    add_rect(s3, Inches(6.90), Inches(4.05), Inches(5.95), Inches(2.55), PANEL2, radius=0.07)
    add_rect(s3, Inches(6.90), Inches(4.05), Inches(0.10), Inches(2.55), OK)
    tb(s3, Inches(7.25), Inches(4.20), Inches(5.35), Inches(0.35),
       "操作员模式", size=16, bold=True, color=OK)
    for j, line in enumerate([
        "运行页：单次触发或连续检测",
        "只暴露业务参数，服务端按范围夹紧",
        "OK / NG、良率、报警；监控可回溯明细",
    ]):
        tb(s3, Inches(7.25), Inches(4.65 + j * 0.42), Inches(5.35), Inches(0.38),
           "▸  " + line, size=13, color=INK)

    prs.save(PPTX_PATH)
    print(f"wrote {PPTX_PATH}")


if __name__ == "__main__":
    build_docx()
    build_pptx()
